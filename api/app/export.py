"""Download my data (task 5.1, SPEC page 4): the user's own rows as Excel, CSV or JSON.

Every query runs in the caller's user_session, so RLS limits it to their rows. Sheets use the
user's units and timezone (units in the headers); JSON keeps the stored units and the ids, the
shape /import reads (task 5.2). Text cells starting with =, +, - or @ get a leading apostrophe
in Excel and CSV so a spreadsheet never runs them as formulas.
"""

import csv
import io
import json
import uuid
import zipfile
from collections.abc import Callable, Iterable
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from typing import IO, Any, Literal, cast
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._write_only import WriteOnlyWorksheet
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.tdee import CM_PER_IN, LB_IN_KG
from app.trend import WEIGHT_ALPHA, Reading, compute_trend

Format = Literal["xlsx", "csv", "json"]
ML_PER_FL_OZ = 29.5735
FORMULA_STARTS = ("=", "+", "-", "@")

# In the order of the dataset picker (SPEC page 4).
DATASETS: dict[str, str] = {
    "food_log": "Food log",
    "daily_totals": "Daily totals",
    "saved_meals": "Saved meals",
    "staples": "Staples",
    "weight": "Weight",
    "measurements": "Measurements",
    "water": "Water",
    "targets": "Targets history",
    "phases": "Phases",
    "checkins": "Weekly check-ins",
    "streaks": "Streaks and badges",
    "custom_foods": "Custom foods",
    "profile": "Profile",
}
# Not limited by the date range: they describe the account, not a day.
UNDATED = {"saved_meals", "staples", "custom_foods", "profile"}


def escape(value: str) -> str:
    """Block spreadsheet formula injection: "=SUM(A1)" becomes "'=SUM(A1)"."""
    return f"'{value}" if value.startswith(FORMULA_STARTS) else value


@dataclass(frozen=True)
class Ctx:
    s: Session
    user_id: uuid.UUID
    tz: str
    imperial: bool
    fl_oz: bool
    date_from: date | None
    date_to: date | None

    def local(self, at: datetime) -> datetime:
        return at.astimezone(ZoneInfo(self.tz)).replace(tzinfo=None)

    def in_range(self, d: date) -> bool:
        return (self.date_from is None or d >= self.date_from) and (
            self.date_to is None or d <= self.date_to
        )

    @property
    def lo(self) -> date:
        return self.date_from or date.min

    @property
    def hi(self) -> date:
        return self.date_to or date.max

    def weight(self, kg: float | Decimal | None) -> float | None:
        if kg is None:
            return None
        return round(float(kg) / LB_IN_KG if self.imperial else float(kg), 2)

    @property
    def weight_unit(self) -> str:
        return "lb" if self.imperial else "kg"

    def water(self, ml: float | int | None) -> float | None:
        if ml is None:
            return None
        return round(float(ml) / ML_PER_FL_OZ, 1) if self.fl_oz else float(ml)

    @property
    def water_unit(self) -> str:
        return "fl oz" if self.fl_oz else "ml"


@dataclass
class Sheet:
    key: str
    title: str
    headers: list[str]
    rows: list[list[Any]]


def _num(v: Any) -> Any:
    return float(v) if isinstance(v, Decimal) else v


def _yes(v: bool | None) -> str | None:
    return None if v is None else ("yes" if v else "no")


def _rows(ctx: Ctx, sql: str, **params: Any) -> list[Any]:
    return list(ctx.s.execute(text(sql), {"u": ctx.user_id, **params}).all())


# --- sheets ----------------------------------------------------------------------------------


def _food_log(ctx: Ctx) -> Sheet:
    rows = _rows(
        ctx,
        "SELECT l.*, m.name AS meal FROM food_logs l LEFT JOIN saved_meals m "
        "ON m.id = l.saved_meal_id WHERE l.user_id = :u AND l.deleted_at IS NULL "
        "AND l.logged_on BETWEEN :a AND :b ORDER BY l.eaten_at, l.id",
        a=ctx.lo,
        b=ctx.hi,
    )
    out = []
    for r in rows:
        at = ctx.local(r.eaten_at)
        out.append(
            [
                r.logged_on,
                at.time().replace(second=0, microsecond=0),
                r.meal,
                r.name,
                r.brand,
                _num(r.quantity),
                r.unit_label,
                _num(r.grams),
                _num(r.kcal),
                _num(r.protein_g),
                _num(r.carbs_g),
                _num(r.fat_g),
                _num(r.fiber_g),
                _num(r.sugar_g),
                _num(r.sodium_mg),
                r.source,
            ]
        )
    return Sheet(
        "food_log",
        DATASETS["food_log"],
        [
            "Date",
            "Time",
            "Saved meal",
            "Food",
            "Brand",
            "Quantity",
            "Unit",
            "Grams",
            "kcal",
            "Protein (g)",
            "Carbs (g)",
            "Fat (g)",
            "Fiber (g)",
            "Sugar (g)",
            "Sodium (mg)",
            "Source",
        ],
        out,
    )


def _targets_by_date(ctx: Ctx) -> list[Any]:
    """Newest first; the row in effect on a day is the first whose effective_from <= day."""
    return _rows(
        ctx,
        "SELECT * FROM targets WHERE user_id = :u AND deleted_at IS NULL "
        "ORDER BY effective_from DESC, created_at DESC, id DESC",
    )


def daily_totals(ctx: Ctx) -> list[dict[str, Any]]:
    """One row per day with food or water logged, in stored units (kcal, g, ml)."""
    food = {
        r.logged_on: r
        for r in _rows(
            ctx,
            "SELECT logged_on, sum(kcal) AS kcal, sum(protein_g) AS protein, "
            "sum(carbs_g) AS carbs, sum(fat_g) AS fat, sum(fiber_g) AS fiber FROM food_logs "
            "WHERE user_id = :u AND deleted_at IS NULL AND logged_on BETWEEN :a AND :b "
            "GROUP BY logged_on",
            a=ctx.lo,
            b=ctx.hi,
        )
    }
    water = {
        r.logged_on: int(r.ml)
        for r in _rows(
            ctx,
            "SELECT logged_on, sum(ml) AS ml FROM water_logs WHERE user_id = :u "
            "AND deleted_at IS NULL AND logged_on BETWEEN :a AND :b GROUP BY logged_on",
            a=ctx.lo,
            b=ctx.hi,
        )
    }
    targets = _targets_by_date(ctx)
    out = []
    for day in sorted(set(food) | set(water)):
        f = food.get(day)
        t = next((t for t in targets if t.effective_from <= day), None)
        on = None
        if f is not None and t is not None:
            on = abs(float(f.kcal) - t.kcal) * 100 <= t.tolerance_pct * t.kcal
        out.append(
            {
                "date": day,
                "kcal": _num(f.kcal) if f else None,
                "protein_g": _num(f.protein) if f else None,
                "carbs_g": _num(f.carbs) if f else None,
                "fat_g": _num(f.fat) if f else None,
                "fiber_g": _num(f.fiber) if f else None,
                "water_ml": water.get(day),
                "target_kcal": t.kcal if t else None,
                "target_protein_g": t.protein_g if t else None,
                "target_carbs_g": t.carbs_g if t else None,
                "target_fat_g": t.fat_g if t else None,
                "on_target": on,
            }
        )
    return out


def _daily_totals(ctx: Ctx) -> Sheet:
    rows = [
        [
            d["date"],
            d["kcal"],
            d["protein_g"],
            d["carbs_g"],
            d["fat_g"],
            d["fiber_g"],
            ctx.water(d["water_ml"]),
            d["target_kcal"],
            d["target_protein_g"],
            d["target_carbs_g"],
            d["target_fat_g"],
            _yes(d["on_target"]),
        ]
        for d in daily_totals(ctx)
    ]
    return Sheet(
        "daily_totals",
        DATASETS["daily_totals"],
        [
            "Date",
            "kcal",
            "Protein (g)",
            "Carbs (g)",
            "Fat (g)",
            "Fiber (g)",
            f"Water ({ctx.water_unit})",
            "Target kcal",
            "Target protein (g)",
            "Target carbs (g)",
            "Target fat (g)",
            "On target",
        ],
        rows,
    )


def _saved_meal_items(ctx: Ctx) -> list[Any]:
    return _rows(
        ctx,
        "SELECT m.id AS meal_id, m.name AS meal, m.use_count, i.*, f.name AS food, "
        "f.kcal_100, f.protein_100, f.carbs_100, f.fat_100 FROM saved_meals m "
        "JOIN saved_meal_items i ON i.saved_meal_id = m.id AND i.deleted_at IS NULL "
        "LEFT JOIN foods f ON f.id = i.food_id WHERE m.user_id = :u AND m.deleted_at IS NULL "
        "ORDER BY m.name, m.id, i.sort_order, i.id",
    )


def _saved_meals(ctx: Ctx) -> Sheet:
    def per(r: Any, col: str) -> float | None:
        v = getattr(r, col)
        return None if v is None else round(float(v) * float(r.grams) / 100, 1)

    rows = [
        [
            r.meal,
            r.food or "(food removed)",
            _num(r.quantity),
            r.unit_label,
            _num(r.grams),
            per(r, "kcal_100"),
            per(r, "protein_100"),
            per(r, "carbs_100"),
            per(r, "fat_100"),
            r.use_count,
        ]
        for r in _saved_meal_items(ctx)
    ]
    return Sheet(
        "saved_meals",
        DATASETS["saved_meals"],
        [
            "Meal",
            "Food",
            "Amount",
            "Unit",
            "Grams",
            "kcal",
            "Protein (g)",
            "Carbs (g)",
            "Fat (g)",
            "Times used",
        ],
        rows,
    )


def _staple_rows(ctx: Ctx) -> list[Any]:
    return _rows(
        ctx,
        "SELECT st.*, coalesce(f.name, m.name) AS what FROM staples st "
        "LEFT JOIN foods f ON f.id = st.food_id LEFT JOIN saved_meals m ON m.id = st.saved_meal_id "
        "WHERE st.user_id = :u AND st.deleted_at IS NULL ORDER BY st.sort_order, st.id",
    )


def _staples(ctx: Ctx) -> Sheet:
    rows = [
        [
            r.what,
            r.unit_label,
            _num(r.unit_grams),
            _num(r.min_units),
            _num(r.max_units),
            _num(r.step_units),
            _yes(r.is_available),
        ]
        for r in _staple_rows(ctx)
    ]
    return Sheet(
        "staples",
        DATASETS["staples"],
        ["Food or meal", "Unit", "Unit grams", "Min", "Max", "Step", "Available"],
        rows,
    )


def _weighins(ctx: Ctx) -> list[Any]:
    return _rows(
        ctx,
        "SELECT * FROM weight_entries WHERE user_id = :u AND deleted_at IS NULL "
        "ORDER BY measured_at, id",
    )


def _weight(ctx: Ctx) -> Sheet:
    entries = _weighins(ctx)
    # The trend runs over every weigh-in, so a range starts from the right value.
    local = [(e, ctx.local(e.measured_at)) for e in entries]
    trend = compute_trend(
        [Reading(at.date().isoformat(), e.measured_at.isoformat(), float(e.kg)) for e, at in local],
        alpha=WEIGHT_ALPHA,
    )
    by_at = {p.at: p for p in trend.points}
    rows = []
    for e, at in local:
        if not ctx.in_range(at.date()):
            continue
        p = by_at.get(e.measured_at.isoformat())
        rows.append(
            [
                at.date(),
                at.time().replace(second=0, microsecond=0),
                ctx.weight(e.kg),
                ctx.weight(p.trend) if p and p.trend is not None else None,
                _yes(bool(p and p.outlier)),
            ]
        )
    unit = ctx.weight_unit
    return Sheet(
        "weight",
        DATASETS["weight"],
        ["Date", "Time", f"Weight ({unit})", f"Trend ({unit})", "Excluded as outlier"],
        rows,
    )


def _measurement_rows(ctx: Ctx) -> list[Any]:
    # Archived metrics are included: archiving hides a metric, it keeps its history.
    return _rows(
        ctx,
        "SELECT e.*, t.name, t.side, t.unit FROM measurement_entries e "
        "JOIN measurement_types t ON t.id = e.type_id WHERE e.user_id = :u "
        "AND e.deleted_at IS NULL AND t.deleted_at IS NULL "
        "AND (e.measured_at AT TIME ZONE :tz)::date BETWEEN :a AND :b "
        "ORDER BY e.measured_at, t.sort_order, e.id",
        tz=ctx.tz,
        a=ctx.lo,
        b=ctx.hi,
    )


def _shown(value: Decimal, unit: str) -> float:
    v = float(value)
    if unit == "in":
        v /= CM_PER_IN
    elif unit == "lb":
        v /= LB_IN_KG
    return round(v, 2)


def _measurements(ctx: Ctx) -> Sheet:
    rows = [
        [
            ctx.local(r.measured_at).date(),
            r.name,
            r.side,
            r.value_text if r.value is None else _shown(r.value, r.unit),
            None if r.unit == "text" else r.unit,
            r.note,
        ]
        for r in _measurement_rows(ctx)
    ]
    return Sheet(
        "measurements",
        DATASETS["measurements"],
        ["Date", "Metric", "Side", "Value", "Unit", "Note"],
        rows,
    )


def _water_rows(ctx: Ctx) -> list[Any]:
    return _rows(
        ctx,
        "SELECT * FROM water_logs WHERE user_id = :u AND deleted_at IS NULL "
        "AND logged_on BETWEEN :a AND :b ORDER BY logged_at, id",
        a=ctx.lo,
        b=ctx.hi,
    )


def _water(ctx: Ctx) -> Sheet:
    rows = []
    for r in _water_rows(ctx):
        at = ctx.local(r.logged_at)
        rows.append([r.logged_on, at.time().replace(second=0, microsecond=0), ctx.water(r.ml)])
    return Sheet("water", DATASETS["water"], ["Date", "Time", f"Amount ({ctx.water_unit})"], rows)


def _target_rows(ctx: Ctx) -> list[Any]:
    """Each change in the range, plus the one already in effect when the range starts."""
    every = sorted(_targets_by_date(ctx), key=lambda t: (t.effective_from, t.created_at, t.id))
    inside = [t for t in every if ctx.in_range(t.effective_from)]
    if ctx.date_from is not None:
        before = [t for t in every if t.effective_from < ctx.date_from]
        if before:
            inside.insert(0, before[-1])
    return inside


def _targets(ctx: Ctx) -> Sheet:
    rows = [
        [
            t.effective_from,
            t.kcal,
            t.protein_g,
            t.carbs_g,
            t.fat_g,
            t.fiber_g,
            ctx.water(t.water_ml),
            t.source,
            t.tdee,
            t.bmr,
        ]
        for t in _target_rows(ctx)
    ]
    return Sheet(
        "targets",
        DATASETS["targets"],
        [
            "Effective from",
            "kcal",
            "Protein (g)",
            "Carbs (g)",
            "Fat (g)",
            "Fiber (g)",
            f"Water ({ctx.water_unit})",
            "Source",
            "TDEE",
            "BMR",
        ],
        rows,
    )


def _phase_rows(ctx: Ctx) -> list[Any]:
    return _rows(
        ctx,
        "SELECT * FROM diet_phases WHERE user_id = :u AND deleted_at IS NULL "
        "AND start_date <= :b AND (end_date IS NULL OR end_date >= :a) ORDER BY start_date",
        a=ctx.lo,
        b=ctx.hi,
    )


def _phases(ctx: Ctx, today: date) -> Sheet:
    from app.routers.phases import summarize

    rows = []
    for p in _phase_rows(ctx):
        summary = None
        if p.start_date <= today:
            last = min(p.end_date, today) if p.end_date else today
            summary = summarize(ctx.s, ctx.user_id, ctx.tz, p.start_date, last)
        change = summary.weight.change_kg if summary and summary.weight else None
        rows.append(
            [
                p.kind,
                p.label,
                p.start_date,
                p.end_date,
                summary.days if summary else None,
                summary.avg_kcal if summary else None,
                ctx.weight(change),
            ]
        )
    return Sheet(
        "phases",
        DATASETS["phases"],
        [
            "Kind",
            "Label",
            "Start",
            "End",
            "Days",
            "Avg kcal",
            f"Weight change ({ctx.weight_unit})",
        ],
        rows,
    )


def _checkin_rows(ctx: Ctx) -> list[Any]:
    return _rows(
        ctx,
        "SELECT * FROM weekly_checkins WHERE user_id = :u AND deleted_at IS NULL "
        "AND week_start + 6 >= :a AND week_start <= :b ORDER BY week_start",
        a=ctx.lo,
        b=ctx.hi,
    )


def _checkins(ctx: Ctx) -> Sheet:
    rows = []
    for r in _checkin_rows(ctx):
        st = r.stats
        w = st.get("weight") or {}
        m = st.get("measurement") or {}
        rows.append(
            [
                r.week_start,
                st.get("days_logged"),
                st.get("avg_kcal"),
                st.get("avg_target_kcal"),
                st.get("days_on_target"),
                st.get("avg_protein_g"),
                st.get("avg_target_protein_g"),
                ctx.weight(w.get("start_kg")),
                ctx.weight(w.get("end_kg")),
                ctx.weight(w.get("change_kg")),
                m.get("name"),
                ctx.water(st.get("water_ml")),
                st.get("fiber_g"),
                (st.get("phase") or {}).get("name"),
            ]
        )
    wu = ctx.weight_unit
    return Sheet(
        "checkins",
        DATASETS["checkins"],
        [
            "Week start",
            "Days logged",
            "Avg kcal",
            "Target kcal",
            "Days on target",
            "Avg protein (g)",
            "Target protein (g)",
            f"Trend start ({wu})",
            f"Trend end ({wu})",
            f"Trend change ({wu})",
            "Fastest-changing measurement",
            f"Avg water ({ctx.water_unit})",
            "Avg fiber (g)",
            "Phase",
        ],
        rows,
    )


STREAK_KINDS = {
    "counted": "counted",
    "freeze_earned": "freeze earned",
    "freeze_used": "freeze used",
    "repaired": "repaired",
    "broken": "broken",
}


def _streak_rows(ctx: Ctx) -> tuple[list[Any], list[Any]]:
    events = _rows(
        ctx,
        "SELECT * FROM streak_events WHERE user_id = :u AND on_date BETWEEN :a AND :b "
        "ORDER BY on_date, id",
        a=ctx.lo,
        b=ctx.hi,
    )
    badges = [
        b
        for b in _rows(ctx, "SELECT * FROM user_badges WHERE user_id = :u ORDER BY earned_at")
        if ctx.in_range(ctx.local(b.earned_at).date())
    ]
    return events, badges


def _streaks(ctx: Ctx) -> Sheet:
    from app.goals import BADGES

    names = {b.key: b.name for b in BADGES}
    events, badges = _streak_rows(ctx)
    rows: list[list[Any]] = [[e.on_date, STREAK_KINDS.get(e.kind, e.kind), None] for e in events]
    rows += [
        [ctx.local(b.earned_at).date(), "badge earned", names.get(b.badge, b.badge)] for b in badges
    ]
    rows.sort(key=lambda r: r[0])
    return Sheet("streaks", DATASETS["streaks"], ["Date", "Kind", "Detail"], rows)


def _custom_food_rows(ctx: Ctx) -> list[Any]:
    return _rows(
        ctx,
        "SELECT f.*, (SELECT coalesce(json_agg(json_build_object('id', su.id, 'label', su.label, "
        "'grams', su.grams, 'is_default', su.is_default) ORDER BY su.is_default DESC, su.label), "
        "'[]'::json) FROM serving_units su WHERE su.food_id = f.id AND su.deleted_at IS NULL) "
        "AS units FROM foods f WHERE f.owner_id = :u AND f.deleted_at IS NULL "
        "ORDER BY f.name, f.id",
    )


def _custom_foods(ctx: Ctx) -> Sheet:
    def units(r: Any) -> str:
        return "; ".join(f"{u['label']} = {float(u['grams']):g} g" for u in r.units)

    rows = [
        [
            r.name,
            r.brand,
            _num(r.kcal_100),
            _num(r.protein_100),
            _num(r.carbs_100),
            _num(r.fat_100),
            _num(r.fiber_100),
            _num(r.sugar_100),
            _num(r.sodium_mg_100),
            units(r) or None,
        ]
        for r in _custom_food_rows(ctx)
    ]
    return Sheet(
        "custom_foods",
        DATASETS["custom_foods"],
        [
            "Name",
            "Brand",
            "kcal per 100 g",
            "Protein per 100 g (g)",
            "Carbs per 100 g (g)",
            "Fat per 100 g (g)",
            "Fiber per 100 g (g)",
            "Sugar per 100 g (g)",
            "Sodium per 100 g (mg)",
            "Serving units",
        ],
        rows,
    )


def _profile_row(ctx: Ctx) -> Any:
    return ctx.s.execute(
        text(
            "SELECT u.email, u.created_at, p.*, st.track_water, st.track_fiber, st.water_unit, "
            "st.gamification, st.streak_reminder FROM users u "
            "JOIN user_profiles p ON p.user_id = u.id "
            "LEFT JOIN user_settings st ON st.user_id = u.id WHERE u.id = :u"
        ),
        {"u": ctx.user_id},
    ).one()


def _profile(ctx: Ctx) -> Sheet:
    p = _profile_row(ctx)
    height = None
    if p.height_cm is not None:
        height = round(float(p.height_cm) / CM_PER_IN, 1) if ctx.imperial else float(p.height_cm)
    rows: list[list[Any]] = [
        ["Name", p.display_name],
        ["Email", p.email],
        ["Account created", ctx.local(p.created_at).date()],
        ["Timezone", p.timezone],
        ["Units", p.unit_system],
        ["Sex", p.sex],
        ["Age", p.age_years],
        [f"Height ({'in' if ctx.imperial else 'cm'})", height],
        [f"Goal weight ({ctx.weight_unit})", ctx.weight(p.goal_weight_kg)],
        ["Theme", p.theme],
        ["Track water", _yes(p.track_water)],
        ["Track fiber", _yes(p.track_fiber)],
        ["Water unit", ctx.water_unit],
        ["Streaks and badges", _yes(p.gamification)],
        ["Streak reminder", p.streak_reminder],
    ]
    return Sheet("profile", DATASETS["profile"], ["Field", "Value"], rows)


SHEETS: dict[str, Callable[[Ctx], Sheet]] = {
    "food_log": _food_log,
    "daily_totals": _daily_totals,
    "saved_meals": _saved_meals,
    "staples": _staples,
    "weight": _weight,
    "measurements": _measurements,
    "water": _water,
    "targets": _targets,
    "checkins": _checkins,
    "streaks": _streaks,
    "custom_foods": _custom_foods,
    "profile": _profile,
}


def build_sheets(ctx: Ctx, keys: Iterable[str], today: date) -> list[Sheet]:
    wanted = set(keys)
    out = []
    for key in DATASETS:
        if key not in wanted:
            continue
        out.append(_phases(ctx, today) if key == "phases" else SHEETS[key](ctx))
    return out


# --- canonical JSON --------------------------------------------------------------------------


def _plain(row: Any, drop: tuple[str, ...] = ("user_id", "owner_id")) -> dict[str, Any]:
    return {k: v for k, v in row._mapping.items() if k not in drop}


def _canonical(ctx: Ctx, key: str, today: date) -> Any:
    if key == "food_log":
        return [
            _plain(r, ("user_id", "meal"))
            for r in _rows(
                ctx,
                "SELECT * FROM food_logs WHERE user_id = :u AND deleted_at IS NULL "
                "AND logged_on BETWEEN :a AND :b ORDER BY eaten_at, id",
                a=ctx.lo,
                b=ctx.hi,
            )
        ]
    if key == "daily_totals":
        return daily_totals(ctx)
    if key == "saved_meals":
        meals: dict[Any, dict[str, Any]] = {}
        for r in _rows(
            ctx,
            "SELECT * FROM saved_meals WHERE user_id = :u AND deleted_at IS NULL ORDER BY name, id",
        ):
            meals[r.id] = {**_plain(r), "items": []}
        for r in _rows(
            ctx,
            "SELECT * FROM saved_meal_items WHERE user_id = :u AND deleted_at IS NULL "
            "ORDER BY sort_order, id",
        ):
            if r.saved_meal_id in meals:
                meals[r.saved_meal_id]["items"].append(_plain(r, ("user_id", "saved_meal_id")))
        return list(meals.values())
    if key == "staples":
        return {
            "staples": [
                _plain(r, ("user_id", "what"))
                for r in _rows(
                    ctx,
                    "SELECT * FROM staples WHERE user_id = :u AND deleted_at IS NULL "
                    "ORDER BY sort_order, id",
                )
            ],
            "never_pairs": [
                _plain(r)
                for r in _rows(
                    ctx,
                    "SELECT * FROM pair_preferences WHERE user_id = :u AND deleted_at IS NULL",
                )
            ],
        }
    if key == "weight":
        return [_plain(r) for r in _weighins(ctx) if ctx.in_range(ctx.local(r.measured_at).date())]
    if key == "measurements":
        return {
            "types": [
                _plain(r)
                for r in _rows(
                    ctx,
                    "SELECT * FROM measurement_types WHERE user_id = :u AND deleted_at IS NULL "
                    "ORDER BY sort_order, id",
                )
            ],
            "entries": [
                _plain(r, ("user_id", "name", "side", "unit")) for r in _measurement_rows(ctx)
            ],
        }
    if key == "water":
        return [_plain(r) for r in _water_rows(ctx)]
    if key == "targets":
        return [_plain(r) for r in _target_rows(ctx)]
    if key == "phases":
        return [_plain(r) for r in _phase_rows(ctx)]
    if key == "checkins":
        return [_plain(r) for r in _checkin_rows(ctx)]
    if key == "streaks":
        events, badges = _streak_rows(ctx)
        return {"events": [_plain(e) for e in events], "badges": [_plain(b) for b in badges]}
    if key == "custom_foods":
        return [_plain(r) for r in _custom_food_rows(ctx)]
    if key == "profile":
        p = _profile_row(ctx)
        keep = (
            "display_name",
            "timezone",
            "unit_system",
            "sex",
            "age_years",
            "age_recorded_on",
            "height_cm",
            "goal_weight_kg",
            "theme",
        )
        settings = ("track_water", "track_fiber", "water_unit", "gamification", "streak_reminder")
        return {
            "email": p.email,
            "profile": {k: getattr(p, k) for k in keep},
            "settings": {k: getattr(p, k) for k in settings},
        }
    raise KeyError(key)


def _count(value: Any) -> int:
    if isinstance(value, list):
        return len(value)
    if isinstance(value, dict) and "profile" in value:
        return 1
    if isinstance(value, dict):
        return sum(len(v) for v in value.values() if isinstance(v, list))
    return 0


def _json_default(v: Any) -> Any:
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, datetime | date | time):
        return v.isoformat()
    if isinstance(v, uuid.UUID | bytes):
        return str(v)
    raise TypeError(type(v).__name__)


# --- writers ---------------------------------------------------------------------------------

DATE_FORMAT = "yyyy-mm-dd"
TIME_FORMAT = "hh:mm"


def _xlsx_value(v: Any) -> Any:
    if isinstance(v, str):
        return escape(v)
    return _num(v)


def _width(v: Any) -> int:
    if isinstance(v, date | time):
        return 11
    return min(len(str(v)), 50) if v is not None else 0


def write_xlsx(sheets: list[Sheet], out: IO[bytes]) -> None:
    wb = Workbook(write_only=True)
    bold = Font(bold=True)
    for sh in sheets:
        ws: WriteOnlyWorksheet = wb.create_sheet(sh.title)
        widths = [len(h) for h in sh.headers]
        for row in sh.rows[:200]:
            for i, v in enumerate(row):
                widths[i] = max(widths[i], _width(v))
        for i, w in enumerate(widths, start=1):
            # Write-only sheets take column widths before the first row (missing from the stubs).
            cast(Any, ws).column_dimensions[get_column_letter(i)].width = w + 2
        ws.freeze_panes = "A2"
        header = []
        for h in sh.headers:
            c = WriteOnlyCell(ws, value=h)
            c.font = bold
            header.append(c)
        ws.append(header)
        for row in sh.rows:
            cells: list[Any] = []
            for v in row:
                if isinstance(v, datetime):
                    c = WriteOnlyCell(ws, value=v)
                    c.number_format = f"{DATE_FORMAT} {TIME_FORMAT}"
                    cells.append(c)
                elif isinstance(v, date):
                    c = WriteOnlyCell(ws, value=cast(Any, v))
                    c.number_format = DATE_FORMAT
                    cells.append(c)
                elif isinstance(v, time):
                    c = WriteOnlyCell(ws, value=cast(Any, v))
                    c.number_format = TIME_FORMAT
                    cells.append(c)
                else:
                    cells.append(_xlsx_value(v))
            ws.append(cells)
    wb.save(out)


def _csv_value(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, time):
        return v.strftime("%H:%M")
    if isinstance(v, date | datetime):
        return v.isoformat()
    if isinstance(v, float):
        return f"{v:g}" if abs(v) < 1e15 else str(v)
    if isinstance(v, str):
        return escape(v)
    return str(_num(v))


def csv_name(sheet: Sheet) -> str:
    return sheet.title.lower().replace(" ", "-") + ".csv"


def write_csv_zip(sheets: list[Sheet], out: IO[bytes]) -> None:
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for sh in sheets:
            buf = io.StringIO()
            # A byte-order mark so Excel opens accents correctly.
            buf.write("﻿")
            w = csv.writer(buf, lineterminator="\r\n")
            w.writerow(sh.headers)
            for row in sh.rows:
                w.writerow([_csv_value(v) for v in row])
            z.writestr(csv_name(sh), buf.getvalue().encode("utf-8"))


def write_json(ctx: Ctx, keys: list[str], today: date, now: datetime, out: IO[bytes]) -> int:
    data = {k: _canonical(ctx, k, today) for k in DATASETS if k in keys}
    doc = {
        "app": "TRIM",
        "format": "trim-export",
        "version": 1,
        "exported_at": now,
        "timezone": ctx.tz,
        "date_from": ctx.date_from,
        "date_to": ctx.date_to,
        "units": "Stored units: kg, cm, g, mg, ml and kcal; times in UTC.",
        "datasets": data,
    }
    out.write(json.dumps(doc, default=_json_default, ensure_ascii=False, indent=1).encode())
    return sum(_count(v) for k, v in data.items() if k != "profile")


def counted_rows(sheets: list[Sheet]) -> int:
    """Rows across the sheets, not counting the profile (it is never empty)."""
    return sum(len(sh.rows) for sh in sheets if sh.key != "profile")


def file_name(fmt: Format, today: date) -> str:
    ext = {"xlsx": "xlsx", "csv": "zip", "json": "json"}[fmt]
    return f"trim-export-{today.isoformat()}.{ext}"


MEDIA_TYPES: dict[str, str] = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "application/zip",
    "json": "application/json",
}

ONE_MINUTE = timedelta(minutes=1)
