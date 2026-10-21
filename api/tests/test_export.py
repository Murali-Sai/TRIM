"""Task 5.1: Download my data. AC: a seeded account's export opened with openpyxl matches the
expected cells, types and units; a food named "=SUM(A1)" exports as text."""

import csv
import io
import json
import zipfile
from datetime import UTC, date, datetime, time, timedelta
from typing import Any, get_args

import pytest
from openpyxl import load_workbook
from sqlalchemy import Engine, text

from app import export as ex
from app.routers.export import Dataset
from tests.conftest import CsrfClient
from tests.fakes import Fakes
from tests.helpers import log_in, sign_up, user_row


@pytest.fixture(autouse=True)
def evening(fakes: Fakes) -> None:
    fakes.clock.now = datetime(2026, 10, 17, 23, 0, tzinfo=UTC)


def seed(owner_engine: Engine, email: str) -> None:
    """New York, imperial: a custom food named =SUM(A1), two logs, weigh-ins, a waist reading
    in inches, water, a target, a phase, a staple and a badge."""
    uid = user_row(owner_engine, email).id
    with owner_engine.begin() as c:

        def q(sql: str, **p: Any) -> Any:
            return c.execute(text(sql), {"u": uid, **p})

        food = q(
            "INSERT INTO foods (id, owner_id, source, name, brand, kcal_100, protein_100, "
            "carbs_100, fat_100, fiber_100, sodium_mg_100) VALUES (gen_random_uuid(), :u, "
            "'custom', '=SUM(A1)', '@brand', 200, 10, 20, 5, 3, 400) RETURNING id"
        ).scalar_one()
        q(
            "INSERT INTO serving_units (id, food_id, label, grams, is_default) VALUES "
            "(gen_random_uuid(), :f, '1 bar', 40, true)",
            f=food,
        )
        # 8:30 pm on 16 Oct in New York is 00:30 UTC on 17 Oct.
        q(
            "INSERT INTO food_logs (id, user_id, food_id, log_group_id, eaten_at, logged_on, "
            "name, brand, grams, quantity, unit_label, kcal, protein_g, carbs_g, fat_g, fiber_g, "
            "sodium_mg, source) VALUES (gen_random_uuid(), :u, :f, gen_random_uuid(), "
            "'2026-10-17T00:30:00Z', '2026-10-16', '=SUM(A1)', '@brand', 80, 2, '1 bar', 160, "
            "8, 16, 4, 2.4, 320, 'search'), (gen_random_uuid(), :u, NULL, gen_random_uuid(), "
            "'2026-09-01T16:00:00Z', '2026-09-01', '-Quick add', NULL, NULL, NULL, NULL, 500, "
            "0, 0, 0, NULL, NULL, 'quick_add')",
            f=food,
        )
        q(
            "INSERT INTO weight_entries (id, user_id, measured_at, kg) VALUES "
            "(gen_random_uuid(), :u, '2026-10-16T11:00:00Z', 80), "
            "(gen_random_uuid(), :u, '2026-10-17T11:00:00Z', 79.5)"
        )
        waist = q(
            "UPDATE measurement_types SET is_active = true, archived_at = now() "
            "WHERE user_id = :u AND name = 'Waist' RETURNING id"
        ).scalar_one()
        q(
            "INSERT INTO measurement_entries (id, user_id, type_id, value, measured_at, note) "
            "VALUES (gen_random_uuid(), :u, :t, 91.44, '2026-10-16T12:00:00Z', '+after run')",
            t=waist,
        )
        q(
            "INSERT INTO water_logs (id, user_id, logged_at, logged_on, ml) VALUES "
            "(gen_random_uuid(), :u, '2026-10-16T14:00:00Z', '2026-10-16', 500)"
        )
        q("DELETE FROM targets WHERE user_id = :u")
        q(
            "INSERT INTO targets (id, user_id, effective_from, kcal, protein_g, carbs_g, fat_g, "
            "source, water_ml) VALUES (gen_random_uuid(), :u, '2026-08-01', 165, 150, 200, 60, "
            "'custom', 2000)"
        )
        q(
            "INSERT INTO diet_phases (id, user_id, kind, start_date, end_date) VALUES "
            "(gen_random_uuid(), :u, 'cut', '2026-10-01', '2026-10-16')"
        )
        q(
            "INSERT INTO staples (id, user_id, food_id, unit_label, unit_grams, step_units, "
            "min_units, max_units) VALUES (gen_random_uuid(), :u, :f, '1 bar', 40, 1, 1, 3)",
            f=food,
        )
        q(
            "INSERT INTO user_badges (id, user_id, badge, earned_at) VALUES "
            "(gen_random_uuid(), :u, 'streak_7', '2026-10-10T15:00:00Z')"
        )


@pytest.fixture
def user(client: CsrfClient, owner_engine: Engine) -> str:
    email = sign_up(client, timezone="America/New_York", unit_system="imperial")
    assert log_in(client, email).status_code == 200
    seed(owner_engine, email)
    return email


def export(client: CsrfClient, fakes: Fakes | None = None, **body: Any) -> Any:
    res = client.post("/api/export", json=body)
    if fakes is not None:
        fakes.clock.now += timedelta(minutes=1, seconds=1)
    return res


def sheet(wb: Any, title: str) -> list[tuple[Any, ...]]:
    return [tuple(c.value for c in row) for row in wb[title].iter_rows()]


def test_datasets_match_the_api_choices() -> None:
    assert list(get_args(Dataset)) == list(ex.DATASETS)


def test_ac_excel_cells_types_and_units(client: CsrfClient, user: str) -> None:
    res = export(client)
    assert res.status_code == 200, res.text
    assert res.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert res.headers["content-disposition"] == (
        'attachment; filename="trim-export-2026-10-17.xlsx"'
    )
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == list(ex.DATASETS.values())

    food = sheet(wb, "Food log")
    assert food[0][:4] == ("Date", "Time", "Saved meal", "Food")
    assert food[0][9:12] == ("Protein (g)", "Carbs (g)", "Fat (g)")
    assert food[0][14] == "Sodium (mg)"
    # Oldest first; the formula-like names are text with a leading apostrophe.
    assert food[2][0] == datetime(2026, 10, 16)  # a real date cell
    assert food[2][1] == time(20, 30)  # local time in New York
    assert food[2][3] == "'=SUM(A1)"
    assert food[2][4] == "'@brand"
    assert food[2][5:9] == (2, "1 bar", 80, 160)
    assert food[1][3] == "'-Quick add"
    ws = wb["Food log"]
    assert ws["A3"].is_date and ws["A3"].number_format == "yyyy-mm-dd"
    assert ws["D3"].data_type == "s"
    assert ws["I3"].data_type == "n"
    assert ws["A1"].font.bold
    assert ws.freeze_panes == "A2"

    weight = sheet(wb, "Weight")
    assert weight[0] == ("Date", "Time", "Weight (lb)", "Trend (lb)", "Excluded as outlier")
    assert weight[1][0] == datetime(2026, 10, 16)
    assert weight[1][1] == time(7, 0)
    assert weight[1][2] == pytest.approx(176.37)
    assert weight[2][4] == "no"

    meas = sheet(wb, "Measurements")
    # Waist is archived and still exported, in inches (its unit for an imperial account).
    assert meas[1] == (datetime(2026, 10, 16), "Waist", None, 36, "in", "'+after run")

    water = sheet(wb, "Water")
    assert water[0] == ("Date", "Time", "Amount (ml)")
    assert water[1] == (datetime(2026, 10, 16), time(10, 0), 500)

    daily = sheet(wb, "Daily totals")
    assert daily[0][0:2] == ("Date", "kcal")
    by_day = {r[0]: r for r in daily[1:]}
    oct16 = by_day[datetime(2026, 10, 16)]
    assert oct16[1] == 160 and oct16[6] == 500 and oct16[7] == 165
    assert oct16[11] == "yes"  # 160 is within 5% of 165
    assert by_day[datetime(2026, 9, 1)][11] == "no"

    targets = sheet(wb, "Targets history")
    assert targets[1][:2] == (datetime(2026, 8, 1), 165)
    assert targets[1][6] == 2000

    phases = sheet(wb, "Phases")
    assert phases[0][-1] == "Weight change (lb)"
    assert phases[1][:5] == ("cut", None, datetime(2026, 10, 1), datetime(2026, 10, 16), 16)

    foods = sheet(wb, "Custom foods")
    assert foods[1][0] == "'=SUM(A1)"
    assert foods[1][2] == 200
    assert foods[1][-1] == "1 bar = 40 g"
    assert sheet(wb, "Staples")[1] == ("'=SUM(A1)", "1 bar", 40, 1, 3, 1, "yes")
    streaks = sheet(wb, "Streaks and badges")
    assert (datetime(2026, 10, 10), "badge earned", "7-day streak") in streaks
    assert sheet(wb, "Weekly check-ins")[1:] == []

    profile = dict(sheet(wb, "Profile")[1:])
    assert profile["Email"] == user
    assert profile["Timezone"] == "America/New_York"
    assert "Password" not in " ".join(str(k) for k in profile)


def test_csv_zip_has_a_bom_and_escapes_formulas(
    client: CsrfClient, user: str, fakes: Fakes
) -> None:
    res = export(client, format="csv", datasets=["food_log", "weight"])
    assert res.status_code == 200, res.text
    assert res.headers["content-type"] == "application/zip"
    assert 'filename="trim-export-2026-10-17.zip"' in res.headers["content-disposition"]
    z = zipfile.ZipFile(io.BytesIO(res.content))
    assert z.namelist() == ["food-log.csv", "weight.csv"]
    raw = z.read("food-log.csv")
    assert raw.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    assert rows[0][0] == "Date"
    assert rows[2][:4] == ["2026-10-16", "20:30", "", "'=SUM(A1)"]


def test_json_keeps_stored_units_and_ids(client: CsrfClient, user: str) -> None:
    res = export(client, format="json")
    assert res.status_code == 200, res.text
    doc = json.loads(res.content)
    assert (doc["format"], doc["version"]) == ("trim-export", 1)
    d = doc["datasets"]
    assert set(d) == set(ex.DATASETS)
    log = next(r for r in d["food_log"] if r["source"] == "search")
    assert log["name"] == "=SUM(A1)"  # JSON is data, not a spreadsheet: no apostrophe
    assert log["eaten_at"].startswith("2026-10-17T00:30:00")
    assert "user_id" not in log and log["id"]
    assert [w["kg"] for w in d["weight"]] == [80.0, 79.5]
    assert d["measurements"]["entries"][0]["value"] == pytest.approx(91.44)
    assert any(t["archived_at"] for t in d["measurements"]["types"])
    assert d["custom_foods"][0]["units"][0]["grams"] == 40
    assert d["profile"]["email"] == user
    assert "password_hash" not in json.dumps(d["profile"])


def test_date_range_limits_dated_datasets_only(client: CsrfClient, user: str) -> None:
    res = export(client, format="json", date_from="2026-10-01", date_to="2026-10-31", datasets=None)
    assert res.status_code == 422  # datasets cannot be null
    res = export(client, format="json", date_from="2026-10-01", date_to="2026-10-31")
    d = json.loads(res.content)["datasets"]
    assert [r["logged_on"] for r in d["food_log"]] == ["2026-10-16"]
    assert len(d["custom_foods"]) == 1  # not limited by the range
    # The target set in August was in effect when the range starts.
    assert [t["effective_from"] for t in d["targets"]] == ["2026-08-01"]


def test_one_export_a_minute(client: CsrfClient, user: str, fakes: Fakes) -> None:
    assert export(client, datasets=["weight"]).status_code == 200
    res = export(client, datasets=["weight"])
    assert res.status_code == 429
    assert res.json()["detail"] == (
        "You can download your data once a minute. Try again in a moment."
    )
    fakes.clock.now += timedelta(seconds=61)
    assert export(client, datasets=["weight"]).status_code == 200


def test_each_export_writes_an_audit_row(
    client: CsrfClient, user: str, owner_engine: Engine
) -> None:
    export(client, format="csv", date_from="2026-10-01", datasets=["weight", "water"])
    uid = user_row(owner_engine, user).id
    with owner_engine.connect() as c:
        row = c.execute(text("SELECT * FROM exports WHERE user_id = :u"), {"u": uid}).one()
        events = c.execute(
            text("SELECT event FROM audit_log WHERE user_id = :u AND event = 'export'"),
            {"u": uid},
        ).all()
    assert (row.format, row.date_from, row.datasets, row.row_count) == (
        "csv",
        date(2026, 10, 1),
        ["weight", "water"],
        3,
    )
    assert len(events) == 1


def test_no_data_in_the_range(client: CsrfClient, user: str, owner_engine: Engine) -> None:
    res = export(client, date_from="2020-01-01", date_to="2020-12-31", datasets=["food_log"])
    assert res.status_code == 404
    assert res.json()["detail"] == "No data in this range."
    with owner_engine.connect() as c:
        n = c.execute(
            text("SELECT count(*) FROM exports WHERE user_id = :u"),
            {"u": user_row(owner_engine, user).id},
        ).scalar_one()
    assert n == 0  # a refused export is not counted, so it does not use up the minute


def test_bad_range_is_rejected(client: CsrfClient, user: str) -> None:
    res = export(client, date_from="2026-10-10", date_to="2026-10-01")
    assert res.status_code == 422
    assert res.json()["detail"] == "The start date is after the end date."


def test_an_unconfirmed_account_can_still_export(client: CsrfClient, fakes: Fakes) -> None:
    email = sign_up(client)
    assert log_in(client, email, keep_signed_in=True).status_code == 200
    # Past the 7-day deadline to confirm the email.
    fakes.clock.advance(days=8)
    assert client.get("/api/weights").status_code == 403
    res = client.post("/api/export", json={"datasets": ["profile"]})
    assert res.status_code == 200, res.text


def test_only_the_requesting_users_rows(
    client: CsrfClient, user: str, owner_engine: Engine
) -> None:
    client.post("/api/auth/logout")
    other = CsrfClient(client.app)
    other.get("/api/auth/csrf")
    email = sign_up(other, timezone="America/New_York")
    assert log_in(other, email).status_code == 200
    res = export(other, format="json")
    assert res.status_code == 200, res.text
    body = res.text
    assert "SUM(A1)" not in body and user not in body
    d = json.loads(body)["datasets"]
    assert d["food_log"] == [] and d["weight"] == []
